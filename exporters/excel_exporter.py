"""
Excel exporter using openpyxl.

Layout (left to right):
  基本信息 Basic Info  |  成本项 Costs  |  收入项 Revenue  |  利润计算 Profit

The sheet has two header rows:
  Row 1 = group bands (merged, color-coded)
  Row 2 = field headers
  Row 3+ = data

Cost / revenue items are ALWAYS listed even when we cannot scrape them, so
the user can fill the blanks in later. Profit and profit margin are live
Excel formulas, so editing any cost/revenue cell recalculates automatically.

Logistics cost auto-estimation
-------------------------------
When package_weight and package_dimensions are available (from SellerSprite),
freight_cost, storage_fee, and other_cost (inbound placement fee) are
pre-filled with estimates so the profit formula is useful from the start.
The user can override any cell; formulas downstream recalculate instantly.

Estimation rules (all per unit):
  头程运费  = package_weight_kg × SEA_FREIGHT_RATE_PER_KG   (default $3.5/kg)
  仓储费    = package_volume_ft³ × STORAGE_RATE_PER_CUFT    (default $0.87/ft³/month)
  入库配置费 = weight tier: ≤0.45 kg→$0.21, ≤0.9→$0.27, else→$0.30
               (Amazon 2024 inbound placement fee, minimal-split tier)
"""
import os
import re
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ── Logistics estimation constants ───────────────────────────────────────────
SEA_FREIGHT_RATE   = 3.5    # USD per kg  (sea freight China→US, typical mid-2025)
STORAGE_RATE       = 0.87   # USD per cubic foot per month  (Amazon Jan-Sep non-peak)
# Amazon 2024 inbound placement fee (minimal-split tier, standard size)
_INBOUND_TIERS = [(0.45, 0.21), (0.90, 0.27), (float("inf"), 0.30)]


def _parse_weight_kg(raw: str) -> float | None:
    """
    Parse a weight string into kilograms.  Handles all formats seen in the wild:
      "2.31 pounds (1.05 kg)"   → 1.05
      "3.17 ounces (89.87 g)"   → 0.08987
      "140 g (140.00 g)"        → 0.140
      "0.22 pounds"             → 0.0998
      "1.05 kg"                 → 1.05
    Priority: parenthetical SI value > explicit unit.
    """
    if not raw:
        return None
    # 1. Parenthetical SI value: "(1.05 kg)" or "(89.87 g)"
    m = re.search(r'\((\d+\.?\d*)\s*(kg|g)\)', raw, re.I)
    if m:
        v, unit = float(m.group(1)), m.group(2).lower()
        return v if unit == "kg" else v / 1000

    # 2. Explicit unit without parentheses
    m = re.search(r'(\d+\.?\d*)\s*(kg|g|pounds?|lbs?|ounces?|oz)', raw, re.I)
    if not m:
        return None
    v, unit = float(m.group(1)), m.group(2).lower()
    if unit in ("kg",):
        return v
    if unit in ("g",):
        return v / 1000
    if unit.startswith("pound") or unit.startswith("lb"):
        return v * 0.453592
    if unit.startswith("ounce") or unit == "oz":
        return v * 0.0283495
    return None


def _parse_dimensions_in(raw: str) -> tuple[float, float, float] | None:
    """
    Parse a dimensions string into (L, W, H) in inches.  Handles:
      "16.9 x 12.3 x 2.2 inches"
      "42.9 x 31.2 x 5.6 cm"
      "16.9 x 12.3 x 2.2"          (assume inches when no unit)
    Returns None when fewer than three numbers are found.
    """
    if not raw:
        return None
    nums = re.findall(r'\d+\.?\d*', raw)
    if len(nums) < 3:
        return None
    l, w, h = float(nums[0]), float(nums[1]), float(nums[2])
    unit_m = re.search(r'\b(inches?|in\b|cm|centimeters?)\b', raw, re.I)
    unit = unit_m.group(1).lower() if unit_m else "in"
    if unit.startswith("cm") or unit.startswith("cent"):
        l, w, h = l / 2.54, w / 2.54, h / 2.54
    return l, w, h


def _estimate_logistics(product: dict) -> dict:
    """
    Return estimated freight_cost, storage_fee, other_cost (inbound placement)
    based on package_weight and package_dimensions.  Returns {} when data is
    insufficient so the caller can skip pre-filling.
    """
    kg = _parse_weight_kg(product.get("package_weight") or product.get("item_weight") or "")
    dims = _parse_dimensions_in(product.get("package_dimensions") or product.get("product_dimensions") or "")

    result = {}
    if kg is not None:
        result["freight_cost"] = round(kg * SEA_FREIGHT_RATE, 2)
        for limit, fee in _INBOUND_TIERS:
            if kg <= limit:
                result["other_cost"] = fee
                break

    if dims is not None:
        l, w, h = dims
        vol_ft3 = (l * w * h) / 1728.0   # 1728 in³ per ft³
        result["storage_fee"] = round(vol_ft3 * STORAGE_RATE, 2)

    return result


# Each column: (key, header, kind, group)
#   kind: "text"  -> raw string
#         "money" -> numeric value parsed from strings like "$54.99"
#         "formula:<name>" -> Excel formula generated per row
#   group: used for the merged band in row 1
G_BASIC   = "基本信息 Basic Info"
G_COST    = "成本项 Costs"
G_REVENUE = "收入项 Revenue"
G_PROFIT  = "利润计算 Profit"
G_AI      = "AI评价 AI Review"

COLUMNS = [
    # ---------------- 基本信息 Basic Info ----------------
    ("asin",                     "ASIN",                          "text",  G_BASIC),
    ("title",                    "产品标题(Title)",                "text",  G_BASIC),
    ("brand",                    "品牌(Brand)",                    "text",  G_BASIC),
    ("seller_name",              "卖家(Seller)",                   "text",  G_BASIC),
    ("sif_seller_count",         "卖家数量(Seller Count)",          "text",  G_BASIC),
    ("bsr",                      "BSR排名(BSR)",                   "text",  G_BASIC),
    ("main_category",            "主类目(Main Category)",          "text",  G_BASIC),
    ("subcategory",              "子类目(Subcategory)",            "text",  G_BASIC),
    ("rating",                   "评分(Rating)",                   "text",  G_BASIC),
    ("review_count",             "评论数(Reviews)",                "text",  G_BASIC),
    ("ss_rating",                "评分_评分数(Rating/Count)",       "text",  G_BASIC),
    ("ss_monthly_sales_parent",  "月销量-父体(Sales Parent)",       "text",  G_BASIC),
    ("ss_monthly_sales_child",   "月销量-子体(Sales Child)",        "text",  G_BASIC),
    ("listing_date",             "上架时间(Launch Date)",          "text",  G_BASIC),
    ("variation_count",          "变体数量(Variations)",            "text",  G_BASIC),
    ("ss_style",                 "款式(Style)",                    "text",  G_BASIC),
    ("item_weight",              "商品重量(Item Weight)",           "text",  G_BASIC),
    ("product_dimensions",       "商品尺寸(Product Dimensions)",    "text",  G_BASIC),
    ("package_weight",           "包裹重量(Package Weight)",        "text",  G_BASIC),
    ("package_dimensions",       "包裹尺寸(Package Dimensions)",    "text",  G_BASIC),
    ("fulfillment",              "配送方式(Fulfillment)",           "text",  G_BASIC),
    ("ss_shipping_days",         "配送时长(Shipping Days)",         "text",  G_BASIC),
    ("coupon_discount",          "优惠券(Coupon)",                 "text",  G_BASIC),
    ("ss_total_traffic",         "全部流量词(Total Traffic)",       "text",  G_BASIC),
    ("ss_organic_traffic",       "自然搜索词(Organic Traffic)",     "text",  G_BASIC),
    ("ss_ad_traffic",            "广告流量词(Ad Traffic)",          "text",  G_BASIC),
    ("ss_recommend_traffic",     "搜索推荐词(Recommend Traffic)",   "text",  G_BASIC),
    ("bullet_points",            "卖点(Bullet Points)",            "text",  G_BASIC),
    ("product_description",      "产品描述(Description)",          "text",  G_BASIC),
    ("customers_say_summary",    "买家综合评价(Customers Say)",     "text",  G_BASIC),
    ("customers_say_topics",     "评价话题标签(Review Topics)",     "text",  G_BASIC),
    ("main_image_url",           "主图URL(Main Image)",            "text",  G_BASIC),
    ("all_image_urls",           "全部图片URL(All Images)",         "text",  G_BASIC),
    ("product_url",              "产品URL(Product URL)",           "text",  G_BASIC),
    ("page_number",              "页码(Page)",                     "text",  G_BASIC),
    ("collection_timestamp",     "采集时间(Timestamp)",            "text",  G_BASIC),

    # ---------------- 成本项 Costs (per unit) ----------------
    # 金额已知类：填金额 → 公式算占比
    # freight_cost / storage_fee / other_cost are auto-estimated from
    # package_weight & package_dimensions; user can overwrite any cell.
    ("product_cost",             "采购成本(Product Cost)",              "money",                      G_COST),
    ("product_cost_pct",         "采购占比%",                           "formula:pct:product_cost",   G_COST),
    ("freight_cost",             "头程运费(Freight)★估",               "money",                      G_COST),
    ("freight_cost_pct",         "头程占比%",                           "formula:pct:freight_cost",   G_COST),
    ("ss_fba_fee",               "FBA配送费(FBA Fee)",                  "money",                      G_COST),
    ("ss_fba_fee_pct",           "FBA占比%",                            "formula:pct:ss_fba_fee",     G_COST),
    ("storage_fee",              "仓储费(Storage Fee)★估",             "money",                      G_COST),
    ("storage_fee_pct",          "仓储占比%",                           "formula:pct:storage_fee",    G_COST),
    ("other_cost",               "入库配置费(Inbound Fee)★估",         "money",                      G_COST),
    ("other_cost_pct",           "入库费占比%",                         "formula:pct:other_cost",     G_COST),
    # 占比已知类：填占比% → 公式算金额
    ("referral_fee_pct",         "平台佣金占比%(填小数如0.15)",      "pct_input:0.15",             G_COST),
    ("referral_fee",             "平台佣金(Referral Fee)",          "formula:from_pct:referral_fee_pct", G_COST),
    ("ad_cost_pct",              "广告费占比%(填小数如0.10)",        "pct_input:0",                G_COST),
    ("ad_cost",                  "广告费(Ad Cost)",                "formula:from_pct:ad_cost_pct", G_COST),
    ("return_cost_pct",          "退货占比%(填小数如0.05)",          "pct_input:0",                G_COST),
    ("return_cost",              "退货成本(Returns)",              "formula:from_pct:return_cost_pct", G_COST),
    # 汇总
    ("total_cost",               "总成本(Total Cost)",             "formula:total_cost",         G_COST),
    ("total_cost_pct",           "总成本占比%",                     "formula:pct:total_cost",     G_COST),

    # ---------------- 收入项 Revenue (per unit) ----------------
    ("price",                    "售价(Selling Price)",            "money", G_REVENUE),
    ("other_revenue",            "其他收入(Other Revenue)",         "money", G_REVENUE),
    ("total_revenue",            "总收入(Total Revenue)",           "formula:total_revenue", G_REVENUE),
    ("ss_monthly_revenue",       "月销售额-参考(Monthly Rev. ref)", "text",  G_REVENUE),

    # ---------------- 利润计算 Profit ----------------
    ("profit",                   "单件利润(Profit)",                "formula:profit",        G_PROFIT),
    ("profit_margin",            "利润率(Profit Margin)",           "formula:profit_margin", G_PROFIT),
    ("ss_gross_margin",          "毛利率-卖家精灵(SS Margin ref)",  "text",  G_PROFIT),

    # ---------------- AI评价 AI Review ----------------
    ("ai_evaluation",            "AI选品评价(AI Evaluation)",       "text",  G_AI),
]

# Styling
HEADER_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
GROUP_FONT  = Font(bold=True, color="FFFFFF", size=12)
GROUP_FILLS = {
    G_BASIC:   PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid"),
    G_COST:    PatternFill(start_color="C0504D", end_color="C0504D", fill_type="solid"),
    G_REVENUE: PatternFill(start_color="4F8A4F", end_color="4F8A4F", fill_type="solid"),
    G_PROFIT:  PatternFill(start_color="E0A526", end_color="E0A526", fill_type="solid"),
    G_AI:      PatternFill(start_color="6F42C1", end_color="6F42C1", fill_type="solid"),
}
ROW_FILL_A  = PatternFill(start_color="EBF3FB", end_color="EBF3FB", fill_type="solid")
ROW_FILL_B  = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
THIN = Side(style="thin", color="D0D0D0")
CELL_BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
MAX_COL_WIDTH = 50
MIN_COL_WIDTH = 10
MONEY_FMT = "#,##0.00"
PCT_FMT = "0.0%"
WRAP_KEYS = {"bullet_points", "product_description", "ai_evaluation", "customers_say_summary"}


def _parse_money(value) -> float | None:
    """Extract a float from strings like '$54.99', '1,234.5'. Returns None if none."""
    if value is None:
        return None
    s = str(value).replace(",", "")
    m = re.search(r"-?\d+(\.\d+)?", s)
    return float(m.group(0)) if m else None


def _col_letters() -> dict:
    """Map each column key to its Excel column letter."""
    return {key: get_column_letter(i) for i, (key, *_rest) in enumerate(COLUMNS, start=1)}


def export(products: list, output_path: str,
           market_summary: str = "", keyword: str = "",
           category_map: dict = None) -> str:
    os.makedirs(os.path.dirname(output_path), exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "Amazon数据采集"

    letters = _col_letters()

    # Cost item keys that contribute real monetary values to the total.
    # Includes: money cols + from_pct formula cols (referral/ad/return).
    # Excludes: pct_input cols, formula:pct cols, and total_cost itself.
    cost_value_keys = [
        k for k, _, kind, grp in COLUMNS
        if grp == G_COST
        and kind not in ("formula:total_cost", "pct_input:0", "pct_input:0.15")
        and not kind.startswith("formula:pct")
        and not kind.startswith("pct_input")
    ]

    # ---------- Row 1: group bands ----------
    # Find contiguous runs of the same group and merge them.
    col = 1
    n = len(COLUMNS)
    while col <= n:
        grp = COLUMNS[col - 1][3]
        start = col
        while col <= n and COLUMNS[col - 1][3] == grp:
            col += 1
        end = col - 1
        c = ws.cell(row=1, column=start, value=grp)
        c.font = GROUP_FONT
        c.fill = GROUP_FILLS.get(grp, HEADER_FILL)
        c.alignment = Alignment(horizontal="center", vertical="center")
        if end > start:
            ws.merge_cells(start_row=1, start_column=start, end_row=1, end_column=end)

    # ---------- Row 2: field headers ----------
    for col_idx, (_key, header, _kind, _grp) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=2, column=col_idx, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = CELL_BORDER

    ws.row_dimensions[1].height = 22
    ws.row_dimensions[2].height = 30

    # ---------- Data rows (start at row 3) ----------
    for offset, product in enumerate(products):
        row_idx = offset + 3
        fill = ROW_FILL_A if offset % 2 == 0 else ROW_FILL_B

        # Pre-fill logistics estimates when the user hasn't supplied values.
        # Only touch cells that are genuinely empty so manual inputs win.
        logistics = _estimate_logistics(product)
        for cost_key, est_val in logistics.items():
            if not product.get(cost_key):
                product[cost_key] = est_val

        for col_idx, (key, _header, kind, _grp) in enumerate(COLUMNS, start=1):
            cell = ws.cell(row=row_idx, column=col_idx)

            if kind == "money":
                cell.value = _parse_money(product.get(key))
                cell.number_format = MONEY_FMT
            elif kind.startswith("pct_input:"):
                # User-editable percentage cell. Pre-fill with default decimal value.
                default_str = kind.split(":", 1)[1]
                default_val = float(default_str) if default_str else None
                cell.value = default_val if default_val else None
                cell.number_format = PCT_FMT
            elif kind.startswith("formula:"):
                name = kind.split(":", 1)[1]
                cell.value = _formula(name, row_idx, letters, cost_value_keys)
                is_pct = (name in ("profit_margin",) or name.startswith("pct:"))
                cell.number_format = PCT_FMT if is_pct else MONEY_FMT
            else:  # text
                cell.value = str(product.get(key, "") or "")

            cell.fill = fill
            cell.border = CELL_BORDER
            wrap = key in WRAP_KEYS
            cell.alignment = Alignment(vertical="top" if wrap else "center", wrap_text=wrap)

    # ---------- Column widths ----------
    last_row = ws.max_row
    for col_idx, (_key, header, _kind, _grp) in enumerate(COLUMNS, start=1):
        col_letter = get_column_letter(col_idx)
        max_len = len(header)
        for r in range(3, last_row + 1):
            v = ws.cell(row=r, column=col_idx).value
            max_len = max(max_len, min(len(str(v if v is not None else "")), MAX_COL_WIDTH))
        ws.column_dimensions[col_letter].width = min(max(max_len + 2, MIN_COL_WIDTH), MAX_COL_WIDTH)

    # Freeze the two header rows
    ws.freeze_panes = "A3"

    # ---------- Optional: 选品结论 summary sheet (keyword-level AI analysis) ----------
    if market_summary:
        s = wb.create_sheet(title="选品结论 Conclusion", index=0)
        s.column_dimensions["A"].width = 100
        title = f"关键词「{keyword}」选品结论" if keyword else "选品结论"
        c = s.cell(row=1, column=1, value=title)
        c.font = Font(bold=True, color="FFFFFF", size=14)
        c.fill = PatternFill(start_color="6F42C1", end_color="6F42C1", fill_type="solid")
        c.alignment = Alignment(horizontal="left", vertical="center")
        s.row_dimensions[1].height = 28
        meta = f"生成时间: {datetime.now().strftime('%Y-%m-%d %H:%M')}  |  竞品数: {len(products)}"
        s.cell(row=2, column=1, value=meta).font = Font(color="888888", size=10)
        body = s.cell(row=4, column=1, value=market_summary)
        body.alignment = Alignment(wrap_text=True, vertical="top", horizontal="left")
        body.font = Font(size=11)
        # Tall row so the wrapped text is visible
        line_est = market_summary.count("\n") + max(1, len(market_summary) // 50)
        s.row_dimensions[4].height = min(600, max(120, line_est * 16))

    # ---------- Optional: 产品分类 category sheet ----------
    if category_map:
        _write_category_sheet(wb, products, category_map, keyword)

    wb.save(output_path)
    return output_path


def _write_category_sheet(wb, products: list, category_map: dict, keyword: str):
    """Add a '产品分类 Category' sheet with products grouped and sorted by BSR."""
    import re as _re

    def _bsr_num(p):
        bsr = str(p.get("bsr", "") or "").replace(",", "")
        m = _re.search(r"\d+", bsr)
        return int(m.group()) if m else 999999

    # Build asin -> product lookup
    asin_map = {p.get("asin", ""): p for p in products}

    # Catch any products the AI didn't assign to a category
    assigned = {a for asins in category_map.values() for a in asins}
    unassigned = [p.get("asin", "") for p in products
                  if p.get("asin", "") and p.get("asin", "") not in assigned]
    if unassigned:
        category_map = dict(category_map)  # don't mutate caller's dict
        category_map["其他 Others"] = unassigned

    cs = wb.create_sheet(title="产品分类 Category")
    cs.column_dimensions["A"].width = 12   # ASIN
    cs.column_dimensions["B"].width = 60   # Title
    cs.column_dimensions["C"].width = 12   # BSR
    cs.column_dimensions["D"].width = 15   # Price
    cs.column_dimensions["E"].width = 18   # Monthly Sales
    cs.column_dimensions["F"].width = 12   # Rating
    cs.column_dimensions["G"].width = 12   # Reviews
    cs.column_dimensions["H"].width = 15   # Seller

    # Header row
    hdr_fill = PatternFill(start_color="2E4057", end_color="2E4057", fill_type="solid")
    hdr_font = Font(bold=True, color="FFFFFF", size=11)
    headers = ["ASIN", "产品标题", "BSR排名", "售价", "月销量(父体)", "评分", "评论数", "卖家"]
    for ci, h in enumerate(headers, 1):
        c = cs.cell(row=1, column=ci, value=h)
        c.font = hdr_font
        c.fill = hdr_fill
        c.alignment = Alignment(horizontal="center", vertical="center")
    cs.row_dimensions[1].height = 22

    cat_fill = PatternFill(start_color="E8EAF6", end_color="E8EAF6", fill_type="solid")
    cat_font = Font(bold=True, size=11, color="3949AB")

    row = 2
    for cat_name, asins in category_map.items():
        # Sort products in this category by BSR ascending
        cat_products = [asin_map[a] for a in asins if a in asin_map]
        cat_products.sort(key=_bsr_num)
        if not cat_products:
            continue

        # Category label row — count actual products found
        c = cs.cell(row=row, column=1, value=f"▶ {cat_name}  ({len(cat_products)} 款)")
        c.font = cat_font
        c.fill = cat_fill
        cs.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(headers))
        cs.row_dimensions[row].height = 20
        row += 1

        alt_fill = PatternFill(start_color="F5F5F5", end_color="F5F5F5", fill_type="solid")
        for i, p in enumerate(cat_products):
            fill = alt_fill if i % 2 == 1 else None
            vals = [
                p.get("asin", ""),
                str(p.get("title", ""))[:100],
                p.get("bsr", ""),
                p.get("price", ""),
                p.get("ss_monthly_sales_parent", ""),
                p.get("rating", ""),
                p.get("review_count", ""),
                p.get("seller_name", ""),
            ]
            for ci, v in enumerate(vals, 1):
                cell = cs.cell(row=row, column=ci, value=v)
                if fill:
                    cell.fill = fill
                if ci == 2:
                    cell.alignment = Alignment(wrap_text=False)
            row += 1

        row += 1  # blank separator between categories

    cs.freeze_panes = "A2"


def _formula(name: str, row: int, letters: dict, cost_value_keys: list) -> str:
    """Build the Excel formula string for a computed column on a given row."""
    rev    = f"{letters['total_revenue']}{row}"
    cost   = f"{letters['total_cost']}{row}"
    profit = f"{letters['profit']}{row}"
    price  = f"{letters['price']}{row}"
    other_rev = f"{letters['other_revenue']}{row}"

    # from_pct:<pct_key> — money = price × pct_input cell; return 0 when empty so total_cost sums cleanly
    if name.startswith("from_pct:"):
        pct_key = name[9:]
        pct_cell = f"{letters[pct_key]}{row}"
        return f'=IF(OR({pct_cell}="",{pct_cell}=0),0,{price}*{pct_cell})'

    if name == "referral_fee":
        # Legacy fallback — column now driven by referral_fee_pct input
        return f"={price}*0.15"

    if name == "total_cost":
        # Sum only the real cost cells (money + referral_fee), skip pct columns.
        # IFERROR handles empty/"" cells so the sum never errors.
        parts = "+".join(f"IFERROR({letters[k]}{row}*1,0)" for k in cost_value_keys)
        return f"={parts}"

    if name == "total_revenue":
        return f"=SUM({price},{other_rev})"

    if name == "profit":
        return f'=IF({rev}=0,"",{rev}-IFERROR({cost}*1,0))'

    if name == "profit_margin":
        return f'=IF(OR({rev}=0,{rev}=""),"",({rev}-IFERROR({cost}*1,0))/{rev})'

    # pct:<cost_key> — percentage of selling price
    if name.startswith("pct:"):
        cost_key = name[4:]
        cost_cell = f"{letters[cost_key]}{row}"
        return f'=IF({price}=0,"",{cost_cell}/{price})'

    return ""


def generate_output_filename(keyword_or_url: str, output_dir: str) -> str:
    """Generate a timestamped output filename for the export."""
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    safe_name = "".join(
        c if c.isalnum() or c in " _-" else "_"
        for c in keyword_or_url[:40]
    ).strip().replace(" ", "_")
    filename = f"amazon_{safe_name}_{timestamp}.xlsx"
    return os.path.join(output_dir, filename)
