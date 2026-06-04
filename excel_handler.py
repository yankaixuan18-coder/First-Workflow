"""
Excel file handler for Amazon product research tool.
Parses uploaded Excel files and creates analysis reports.
"""

import io
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


# Mapping from column letters to field names
COLUMN_FIELD_MAP = {
    "B": "product_name",
    "C": "brand",
    "D": "top_asin",
    "E": "top_asin_monthly_sales",
    "F": "parent_monthly_sales",
    "G": "variants",
    "H": "competitor_relevance",
    "I": "seller_origin",
    "J": "annual_feedback",
    "K": "rating",
    "L": "reviews",
    "M": "launch_date",
    "N": "transaction_price",
    "O": "coupon",
    "P": "product_cost_rmb",
    "Q": "historical_price",
    "R": "platform_commission",
    "S": "fba",
    "T": "package_weight",
    "U": "volume_weight",
    "V": "package_length",
    "W": "package_width",
    "X": "package_height",
    "Y": "gross_profit",
    "Z": "gross_margin",
    "AA": "category_node",
    "AB": "keyword1",
    "AC": "trend1",
    "AD": "keyword2",
    "AE": "trend2",
    "AF": "keyword3",
    "AG": "trend3",
    "AH": "keyword4",
    "AI": "trend4",
    "AJ": "keyword5",
    "AK": "trend5",
    "AL": "top_asin_trend",
    "AM": "conversion_rate",
    "AN": "cpc",
    "AO": "cpc_screenshot",
    "AP": "ad_roi",
    "AQ": "visual_images",
    "AR": "visual_video",
    "AS": "visual_aplus",
    "AT": "consumer_profile",
    "AU": "use_cases",
    "AV": "unmet_needs",
    "AW": "pros",
    "AX": "opportunity_pros",
    "AY": "cons",
    "AZ": "opportunity_cons",
    "BA": "purchase_motivation",
    "BB": "feature_ratings",
    "BC": "return_reasons",
    "BD": "repurchase_cycle",
    "BE": "title",
    "BF": "selling_points",
    "BG": "ad_share",
    "BH": "concentration",
    "BI": "ads",
    "BJ": "off_season_storage",
    "BK": "first_leg_weight",
    "BL": "first_leg_size",
    "BM": "first_leg_shipping",
    "BN": "product_cost_usd",
    "BO": "return_rate",
    "BP": "exchange_rate",
    "BQ": "current_first_leg_price",
}


def col_letter_to_index(col_letter: str) -> int:
    """Convert column letter (e.g. 'A', 'AA') to 1-based index."""
    result = 0
    for char in col_letter.upper():
        result = result * 26 + (ord(char) - ord('A') + 1)
    return result


def parse_excel(file_path: str) -> dict:
    """
    Parse an Excel file and extract product data from row 3.

    Args:
        file_path: Path to the Excel file

    Returns:
        Dictionary mapping field names to values
    """
    wb = load_workbook(file_path, data_only=True)
    ws = wb.active

    data = {}
    data_row = 3  # Data is in row 3

    for col_letter, field_name in COLUMN_FIELD_MAP.items():
        col_idx = col_letter_to_index(col_letter)
        cell = ws.cell(row=data_row, column=col_idx)
        value = cell.value

        if value is not None:
            # Convert to string for text fields, keep numbers as numbers
            if isinstance(value, (int, float)):
                data[field_name] = value
            else:
                data[field_name] = str(value).strip()
        else:
            data[field_name] = ""

    return data


def create_report(original_path: str, data: dict, analysis: dict) -> bytes:
    """
    Create an Excel report with original data plus AI analysis in a new sheet.

    Args:
        original_path: Path to the original Excel file (or None if manual entry)
        data: Product data dictionary
        analysis: Analysis results from Claude

    Returns:
        Excel file as bytes
    """
    if original_path:
        wb = load_workbook(original_path)
    else:
        wb = Workbook()
        # Create a basic data sheet
        ws_data = wb.active
        ws_data.title = "产品数据"
        _write_data_sheet(ws_data, data)

    # Create AI analysis sheet
    if "AI分析" in wb.sheetnames:
        del wb["AI分析"]

    ws_analysis = wb.create_sheet("AI分析")
    _write_analysis_sheet(ws_analysis, data, analysis)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()


def _write_data_sheet(ws, data: dict):
    """Write product data to a worksheet."""
    ws.append(["字段", "值"])
    ws["A1"].font = Font(bold=True)
    ws["B1"].font = Font(bold=True)

    field_labels = {
        "product_name": "产品名称",
        "brand": "品牌",
        "top_asin": "热销ASIN",
        "top_asin_monthly_sales": "热销ASIN月销量",
        "parent_monthly_sales": "父体月销量",
        "variants": "变体数量",
        "rating": "评分",
        "reviews": "评论数",
        "gross_profit": "毛利润",
        "gross_margin": "毛润率",
        "transaction_price": "成交金额",
    }

    for field, label in field_labels.items():
        value = data.get(field, "")
        if value:
            ws.append([label, value])


def _write_analysis_sheet(ws, data: dict, analysis: dict):
    """Write AI analysis results to a worksheet."""
    # Styles
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True, size=12)
    section_fill = PatternFill(start_color="2E75B6", end_color="2E75B6", fill_type="solid")
    section_font = Font(color="FFFFFF", bold=True)
    go_fill = PatternFill(start_color="70AD47", end_color="70AD47", fill_type="solid")
    maybe_fill = PatternFill(start_color="FFC000", end_color="FFC000", fill_type="solid")
    nogo_fill = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")
    alt_fill = PatternFill(start_color="DEEAF1", end_color="DEEAF1", fill_type="solid")
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    # Set column widths
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 55
    ws.column_dimensions["C"].width = 18

    row = 1

    # Title
    ws.merge_cells(f"A{row}:C{row}")
    title_cell = ws[f"A{row}"]
    title_cell.value = "亚马逊产品AI选品分析报告"
    title_cell.font = Font(bold=True, size=16, color="FFFFFF")
    title_cell.fill = header_fill
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[row].height = 32
    row += 1

    # Product name
    product_name = data.get("product_name", "未知产品")
    ws.merge_cells(f"A{row}:C{row}")
    cell = ws[f"A{row}"]
    cell.value = f"产品：{product_name}"
    cell.font = Font(bold=True, size=11)
    cell.alignment = Alignment(horizontal="center")
    row += 2

    # Verdict section
    verdict = analysis.get("verdict", "Unknown")
    overall = analysis.get("overall_score", 0)

    ws.merge_cells(f"A{row}:C{row}")
    verdict_cell = ws[f"A{row}"]
    verdict_cell.value = f"综合评分：{overall:.1f} / 10   |   选品建议：{verdict}"
    verdict_cell.font = Font(bold=True, size=14, color="FFFFFF")
    verdict_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[row].height = 30

    if verdict == "Go":
        verdict_cell.fill = go_fill
    elif verdict == "Maybe":
        verdict_cell.fill = maybe_fill
    else:
        verdict_cell.fill = nogo_fill
    row += 2

    # Score dimensions
    ws.merge_cells(f"A{row}:C{row}")
    sec_cell = ws[f"A{row}"]
    sec_cell.value = "各维度评分"
    sec_cell.font = section_font
    sec_cell.fill = section_fill
    sec_cell.alignment = Alignment(horizontal="center")
    row += 1

    # Header row
    for col, header in [("A", "评估维度"), ("B", "说明"), ("C", "得分 (0-10)")]:
        c = ws[f"{col}{row}"]
        c.value = header
        c.font = Font(bold=True)
        c.fill = PatternFill(start_color="BDD7EE", end_color="BDD7EE", fill_type="solid")
        c.border = thin_border
        c.alignment = Alignment(horizontal="center")
    row += 1

    dimensions = [
        ("market_opportunity", "市场机会", "基于月销量、趋势、未被满足需求"),
        ("profitability", "盈利能力", "基于毛利润、毛润率、FBA及头程成本"),
        ("competition_level", "竞争程度", "分数越高竞争越少（对卖家越有利）"),
        ("product_quality_signal", "产品质量信号", "基于评分、优缺点、退货率"),
        ("trend_longevity", "趋势与持续性", "基于近3年关键词及ASIN趋势"),
    ]

    for i, (key, label, desc) in enumerate(dimensions):
        score = analysis.get(key, 0)
        fill = alt_fill if i % 2 == 0 else PatternFill(fill_type=None)

        a_cell = ws[f"A{row}"]
        a_cell.value = label
        a_cell.font = Font(bold=True)
        a_cell.fill = fill
        a_cell.border = thin_border

        b_cell = ws[f"B{row}"]
        b_cell.value = desc
        b_cell.fill = fill
        b_cell.border = thin_border

        c_cell = ws[f"C{row}"]
        c_cell.value = score
        c_cell.fill = fill
        c_cell.border = thin_border
        c_cell.alignment = Alignment(horizontal="center")
        row += 1

    row += 1

    # Reasoning
    ws.merge_cells(f"A{row}:C{row}")
    sec_cell = ws[f"A{row}"]
    sec_cell.value = "分析详情"
    sec_cell.font = section_font
    sec_cell.fill = section_fill
    sec_cell.alignment = Alignment(horizontal="center")
    row += 1

    reasoning = analysis.get("reasoning", "")
    ws.merge_cells(f"A{row}:C{row}")
    r_cell = ws[f"A{row}"]
    r_cell.value = reasoning
    r_cell.alignment = Alignment(wrap_text=True, vertical="top")
    # Estimate height
    lines = max(len(reasoning) // 80 + reasoning.count('\n') + 1, 3)
    ws.row_dimensions[row].height = min(lines * 15, 300)
    row += 2

    # Key risks
    ws.merge_cells(f"A{row}:C{row}")
    sec_cell = ws[f"A{row}"]
    sec_cell.value = "主要风险"
    sec_cell.font = section_font
    sec_cell.fill = PatternFill(start_color="C00000", end_color="C00000", fill_type="solid")
    sec_cell.alignment = Alignment(horizontal="center")
    row += 1

    for risk in analysis.get("key_risks", []):
        ws.merge_cells(f"A{row}:C{row}")
        c = ws[f"A{row}"]
        c.value = f"• {risk}"
        c.alignment = Alignment(wrap_text=True)
        ws.row_dimensions[row].height = max(int(len(risk) / 60) * 15 + 15, 18)
        row += 1

    row += 1

    # Key opportunities
    ws.merge_cells(f"A{row}:C{row}")
    sec_cell = ws[f"A{row}"]
    sec_cell.value = "主要机会"
    sec_cell.font = section_font
    sec_cell.fill = PatternFill(start_color="375623", end_color="375623", fill_type="solid")
    sec_cell.alignment = Alignment(horizontal="center")
    row += 1

    for opp in analysis.get("key_opportunities", []):
        ws.merge_cells(f"A{row}:C{row}")
        c = ws[f"A{row}"]
        c.value = f"• {opp}"
        c.alignment = Alignment(wrap_text=True)
        ws.row_dimensions[row].height = max(int(len(opp) / 60) * 15 + 15, 18)
        row += 1
